from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

out=Path(__file__).parent
source=next(Path('/workspace/attachments').rglob('*.csv'))
raw=pd.read_csv(source,encoding='utf-8-sig')
df=raw.loc[raw['参加意向'].notna()].copy()
target=df.loc[df['参加意向'].eq('内容によっては参加したい')].copy()

def flags(s):
    return s.fillna('').map(lambda v: ','.join(x.strip() for x in v.split(',') if x.strip())).str.get_dummies(sep=',')

def table(col):
    f=flags(target[col])
    return pd.DataFrame({'項目':f.columns,'選択人数':f.sum().values,'選択率（%）':(f.mean()*100).round(1).values}).sort_values('選択人数',ascending=False).reset_index(drop=True)

plans=table('参加インセンティブ_最大3つ選択')
media=table('情報源')
age=target['Age group'].value_counts().reindex(['10代','20代','30代','40代','50代','60代','70代','80代以上'],fill_value=0).rename_axis('年代').reset_index(name='人数')
age['構成比（%）']=(age['人数']/len(target)*100).round(1)
intent=df['参加意向'].value_counts().rename_axis('参加意向').reset_index(name='人数')
intent['構成比（%）']=(intent['人数']/len(df)*100).round(1)
choice_counts=target['参加インセンティブ_最大3つ選択'].str.split(',').str.len()
notes=[
('分析対象','参加意向が「内容によっては参加したい」の回答者を、本資料では潜在層と定義。環境関心の低さを意味しない。'),
('元データ',source.name),
('有効回答',f'{len(df)}件。参加意向が空欄の末尾1行を除外。'),
('潜在層',f'{len(target)}人（有効回答の{len(target)/len(df)*100:.1f}%）。'),
('集計方法','複数回答をカンマで分割し、項目ごとに選択人数を集計。選択率の分母は潜在層472人。合計は100%を超える。'),
('年代','20代以下は少数のため若者の傾向の断定は避ける。学生区分はなく、年代から学生とは判断しない。'),
('選択数の注意',f'最大3つ選択の設問に4つ以上の回答が潜在層で{int((choice_counts>3).sum())}件。主集計は削除せず保持。3つ以内に限定した参考集計を別シートに収録。元フォーム確認が必要。'),
('解釈','参加を促す効果や広報の集客効果を証明するものではない。回答者の選好と情報源の自己申告の集計。'),
('分析範囲','環境配慮行動の実施状況、心理的ハードルを直接測る設問は、このCSVには含まれない。'),
('個別回答','本資料には個別回答や自由記述を掲載せず、集計のみを収録。'),
]
small=target.loc[choice_counts.le(3)]
f=flags(small['参加インセンティブ_最大3つ選択'])
sensitivity=pd.DataFrame({'項目':f.columns,'選択人数':f.sum().values,'選択率（%）':(f.mean()*100).round(1).values,'分母（人）':len(small)}).sort_values('選択人数',ascending=False)
# 年代別も分母を明記。各項目の列が年代。
by_age={}
for col,label in [('参加インセンティブ_最大3つ選択','企画_年代別'),('情報源','情報源_年代別')]:
    f=flags(target[col]); result=pd.DataFrame({'項目':f.columns})
    for a in age['年代']:
        mask=target['Age group'].eq(a); n=int(mask.sum())
        result[f'{a}（n={n}）選択率%']=(f.loc[mask].mean()*100).round(1).values if n else float('nan')
    by_age[label]=result
sheets={'説明・注意点':pd.DataFrame(notes,columns=['項目','内容']),'参加意向_全体':intent,'潜在層_年代構成':age,'潜在層_企画':plans,'潜在層_情報源':media,**by_age,'参考_3つ以内の企画':sensitivity}
with pd.ExcelWriter(out/'潜在層アンケート分析.xlsx',engine='openpyxl') as writer:
    for name,t in sheets.items():
        t.to_excel(writer,sheet_name=name,index=False)
        ws=writer.sheets[name]; ws.freeze_panes='B2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='245C63')
        for row in ws.iter_rows(min_row=2):
            for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
        for i,c in enumerate(t.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=75 if (i==1 and name not in ['説明・注意点','潜在層_年代構成','参加意向_全体']) else (105 if name=='説明・注意点' and i==2 else 25)
        if name=='説明・注意点':
            for i in range(2,len(t)+2): ws.row_dimensions[i].height=48
for name,t in [('潜在層_企画',plans),('潜在層_情報源',media),('潜在層_年代構成',age)]:
    t.to_csv(out/(name+'.csv'),index=False,encoding='utf-8-sig')

def md(t):
    return '| '+' | '.join(t.columns)+' |\n|'+'|'.join(['---']*len(t.columns))+'|\n'+'\n'.join('| '+' | '.join(str(v) for v in row)+' |' for row in t.itertuples(index=False,name=None))
young=int(age.loc[age['年代'].isin(['10代','20代']),'人数'].sum())
text=f'''# ゼロカーボン大作戦：潜在層の参加ニーズ・情報源分析

## 目的と対象
「内容によっては参加したい」と回答した人を「潜在層」と定義し、回遊促進施策と広報施策の具体化に役立てる。これは環境関心の低い層を直接識別したものではない。

有効回答{len(df)}人のうち潜在層は{len(target)}人（{len(target)/len(df)*100:.1f}%）。元CSVの末尾にある参加意向空欄の1行は除外した。

## 潜在層の年代構成

{md(age)}

潜在層の20代以下は{young}人。学生かどうかは識別できず、若者・学生についての一般化は避ける。主に今回回答した大人世代のニーズを把握する資料として用いる。

## 参加したくなる企画

{md(plans)}

特典が最多。短時間、飲食、1人でも楽しめることも上位であり、回遊施策では「短時間で完了できる」「個人で参加できる」「達成に応じて特典を得られる」といった設計を検討できる。ラリーと特典を組み合わせた効果そのものを検証した結果ではない。

## 情報源

{md(media)}

上位媒体を大人世代への広報先候補として検討する。自治体の広報誌とホームページ、テレビと雑誌などは同一選択肢なので個別の効果は比較できない。情報源の選択率は広告の集客効果や閲覧頻度とは異なる。

## 活動・企業共有への活用案

- 回遊施策：少数ブースでも完了できるコース、個人参加、短時間の体験、特典を検討する。
- 広報施策：情報源上位の媒体を中心に、体験時間・参加方法・内容を具体的に告知する。
- 出展企業との共有：短時間・1人でも参加しやすい体験と、暮らしにつながる説明を提案する。集計結果と提案を区別して伝える。
- 実施後：参加者年代、来場のきっかけ、回遊ブース数、体験の感想を確認し、施策を評価する。

## 集計の注意点

選択率は各項目の選択人数÷潜在層{len(target)}人。複数回答なので合計は100%を超える。

「最大3つ選択」の企画設問に4つ以上選んだ回答が潜在層で{int((choice_counts>3).sum())}件ある。主集計では保持し、Excelに3つ以内の回答者{len(small)}人に限った参考集計を収録した。元フォームや転記方法の確認が必要で、参考集計も正解とみなすものではない。

環境配慮行動の実施状況や心理的ハードルは本CSVでは直接分析できない。実際の参加や行動変容の因果効果はこの集計からは判断できない。

## 収録ファイル

- 潜在層アンケート分析.xlsx：説明、全体の参加意向、潜在層年代構成、企画・情報源ランキング、年代別比較、選択数に関する参考集計。
- 潜在層_企画.csv／潜在層_情報源.csv／潜在層_年代構成.csv：表の再利用用。
- analyze.py：元CSVから資料を再生成するPythonスクリプト（pandas、openpyxlが必要）。
'''
(out/'潜在層分析レポート.md').write_text(text,encoding='utf-8')
print('有効回答:',len(df),'潜在層:',len(target),'潜在層20代以下:',young,'4つ以上:',int((choice_counts>3).sum()))
print(plans.head().to_string(index=False)); print(media.head().to_string(index=False))
